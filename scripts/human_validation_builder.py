"""
Human Validation Builder Script
Reads golden_dataset.json and builds human_review.xlsx with native Excel DataValidation dropdowns
and human_review.csv for Stage 2 Human Validation Framework.
"""

import os
import sys
import json
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from app.core.logging_config import setup_logger

logger = setup_logger("human_validation_builder")

REVIEW_OPTIONS = [
    "✅ Correct",
    "⚠ Partially Correct",
    "❌ Wrong Answer",
    "❌ Hallucination",
    "❌ Wrong Source",
    "❌ Needs Better Retrieval",
    "❌ Missing Information"
]


def build_human_review_files():
    golden_json_path = BASE_DIR / "data" / "datasets" / "golden_dataset.json"
    if not golden_json_path.exists():
        logger.error(f"Golden dataset not found at {golden_json_path}")
        sys.exit(1)

    with open(golden_json_path, "r", encoding="utf-8") as f:
        golden_data = json.load(f)

    logger.info(f"Loaded {len(golden_data)} items from golden_dataset.json")

    rows = []
    for item in golden_data:
        cov_status = item.get("coverage_status", "Missing")
        
        # Determine smart initial default for Human Review column
        if cov_status == "Covered":
            default_review = "✅ Correct"
        elif cov_status == "Partial":
            default_review = "⚠ Partially Correct"
        else:
            default_review = "❌ Missing Information"

        row = {
            "Question ID": item.get("question_id") or item.get("id"),
            "Category": item.get("category", "General"),
            "Question": item.get("question", ""),
            "Generated Answer": item.get("golden_answer", ""),
            "Golden Answer": item.get("golden_answer", "") if cov_status == "Covered" else "Factual Ground Truth Required",
            "Source Document": item.get("source_document", "N/A"),
            "Source Chunk": item.get("chunk_id", "N/A"),
            "Confidence Score": item.get("confidence_score", 0.0),
            "Similarity Score": item.get("similarity_score", 0.0),
            "Coverage Status": cov_status,
            "Human Review": default_review,
            "Reviewer Notes": item.get("reason_if_missing", "")
        }
        rows.append(row)

    df = pd.DataFrame(rows)

    # 1. Export CSV
    csv_path = BASE_DIR / "data" / "review" / "human_review.csv"
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(csv_path, index=False, encoding="utf-8")
    logger.info(f"Saved human_review.csv ({len(df)} rows)")

    # 2. Build formatted Excel with native DataValidation Dropdown
    xlsx_path = BASE_DIR / "data" / "review" / "human_review.xlsx"
    df.to_excel(xlsx_path, index=False, engine="openpyxl")

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    # Configure Data Validation for Human Review column (Column K)
    formula_str = '"' + ",".join(REVIEW_OPTIONS) + '"'
    dv = DataValidation(
        type="list",
        formula1=formula_str,
        allow_blank=True,
        showDropDown=False,  # Shows dropdown arrow on cell selection in Excel
        errorTitle="Invalid Selection",
        error="Please select a valid review option from the dropdown menu."
    )

    ws.add_data_validation(dv)
    
    # Add validation range for Human Review column (K2 to K{total_rows+1})
    max_row = len(rows) + 1
    dv.add(f"K2:K{max_row}")

    # Apply professional styling
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    
    # Header styling
    for col_num in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Row styling and column width adjustments
    col_widths = {
        "A": 12,  # Question ID
        "B": 22,  # Category
        "C": 35,  # Question
        "D": 45,  # Generated Answer
        "E": 45,  # Golden Answer
        "F": 24,  # Source Document
        "G": 15,  # Source Chunk
        "H": 16,  # Confidence Score
        "I": 16,  # Similarity Score
        "J": 16,  # Coverage Status
        "K": 26,  # Human Review (Dropdown)
        "L": 35   # Reviewer Notes
    }

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    for row_num in range(2, max_row + 1):
        for col_num in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = thin_border
            
            # Wrap text for text-heavy columns
            if col_num in (3, 4, 5, 12):
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top")

    ws.row_dimensions[1].height = 28
    wb.save(xlsx_path)
    logger.info(f"Saved human_review.xlsx with DataValidation dropdowns on Column K ({max_row-1} items)")


if __name__ == "__main__":
    build_human_review_files()
